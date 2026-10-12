import io
import json
import re
from collections import defaultdict
from pathlib import Path
from xml.sax.saxutils import escape

import numpy as np
import pandas as pd
from PIL import Image as PILImage
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Image as RLImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROWS = 5  # rows in room-wise sheet
SIGN_FILE = Path(__file__).with_name("signature.png")  # Chief Superintendent signature + stamp (room wise PDF)
ROOM_CFG_FILE = Path(__file__).with_name("rooms_config.json")  # remembered room names + default capacity


def load_room_cfg():
    try:
        d = json.loads(ROOM_CFG_FILE.read_text(encoding="utf-8"))
        return {"default_capacity": int(d.get("default_capacity", 25)),
                "room_names": [str(x) for x in d.get("room_names", [])]}
    except Exception:
        return {"default_capacity": 25, "room_names": []}


def save_room_cfg(cfg):
    try:
        ROOM_CFG_FILE.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass  # read-only disk: app still works, it just won't remember


def clean_signature(raw):
    """Any scan/photo of the signature -> PNG whose white background is transparent."""
    img = PILImage.open(io.BytesIO(raw)).convert("RGB")
    a = np.asarray(img).astype(int)
    alpha = np.clip((255 - a.min(axis=2) - 12) * 4, 0, 255).astype("uint8")
    out = np.dstack([a.astype("uint8"), alpha])
    buf = io.BytesIO()
    PILImage.fromarray(out, "RGBA").save(buf, format="PNG")
    return buf.getvalue()


# ---------------------------------------------------------------- logic
def compress(regs):
    """['511524104064','511524104065','511524104067'] -> '511524104064 to 511524104065, 511524104067'"""
    nums = sorted(int(r) for r in regs)
    parts, start, prev = [], nums[0], nums[0]
    for n in nums[1:]:
        if n == prev + 1:
            prev = n
            continue
        parts.append(f"{start} to {prev}" if start != prev else str(start))
        start = prev = n
    parts.append(f"{start} to {prev}" if start != prev else str(start))
    return ", ".join(parts)


def allocate(students, rooms):
    """students: DataFrame(reg_no, department, subcode, ...); rooms: list of (room_no, capacity).
    A "group" is (department, subject code). Each room is filled with 2-3 different groups
    (largest remaining first), preferring groups that write a different subject, so students
    of the same subject are not all together. Inside a room, seats are ordered so that the
    same subject is not above/below or left/right of each other (as far as possible)."""
    students = students.sort_values("reg_no").reset_index(drop=True)
    students["_g"] = students["department"] + " | " + students["subcode"]
    sub_of = dict(zip(students["_g"], students["subcode"]))
    queues = {g: list(x.index) for g, x in students.groupby("_g")}
    placed = []  # (room_no, student_index)
    used_rooms = []
    for room, cap in rooms:
        if not any(queues.values()):
            break
        left, used, used_subs = cap, set(), set()
        while left > 0:
            avail = [g for g in queues if queues[g]]
            if not avail:
                break
            cand = ([g for g in avail if sub_of[g] not in used_subs]
                    or [g for g in avail if g not in used] or avail)
            g = max(cand, key=lambda x: len(queues[x]))
            others = ([x for x in avail if x not in used and sub_of[x] != sub_of[g]]
                      or [x for x in avail if x != g and x not in used])
            n = (left + 1) // 2 if (others and not used) else left
            n = min(n, len(queues[g]))
            for idx in queues[g][:n]:
                placed.append((room, idx))
            queues[g] = queues[g][n:]
            used.add(g)
            used_subs.add(sub_of[g])
            left -= n
        used_rooms.append(room)
    if any(queues.values()):
        raise ValueError("Rooms podhala! Innum rooms add pannunga.")
    # seat order inside each room (column-wise numbering, see room_grid):
    # pos-1 is the seat above, pos-ROWS is the seat to the left
    final = []
    for room in used_rooms:
        groups = defaultdict(list)
        for r, i in placed:
            if r == room:
                groups[students.loc[i, "_g"]].append(i)
        seat_subs, order = [], []
        total = sum(len(v) for v in groups.values())
        for pos in range(total):
            banned = set()
            if pos % ROWS != 0:
                banned.add(seat_subs[pos - 1])
            if pos >= ROWS:
                banned.add(seat_subs[pos - ROWS])
            avail = [g for g in groups if groups[g]]
            ok = [g for g in avail if sub_of[g] not in banned] or avail
            g = max(ok, key=lambda x: len(groups[x]))
            seat_subs.append(sub_of[g])
            order.append(groups[g].pop(0))
        final += [(room, i) for i in order]
    alloc = students.loc[[i for _, i in final]].drop(columns="_g").copy()
    alloc["room_no"] = [r for r, _ in final]
    return alloc.reset_index(drop=True), used_rooms


ALT_MAX = 15  # single-subject room with <= 15 students: seat in rows 1,3,5 only (one empty row between)


def is_alt_room(sub):
    """sub: students of one room. Only one subject can't be mixed, so leave a row between students."""
    return sub.subcode.nunique() == 1 and len(sub) <= ALT_MAX


def seat_layout(n, alt=False):
    """(cols, [(row, col) for every student in seat order]). Numbered column-wise.
    alt=True -> only rows 1,3,5 are used."""
    if alt:
        return -(-n // 3), [(2 * (i % 3), i // 3) for i in range(n)]
    return -(-n // ROWS), [(i % ROWS, i // ROWS) for i in range(n)]


def room_grid(regs, alt=False):
    """Excel room sheet: [S.No, Reg No] pairs, 5 rows."""
    cols, pos = seat_layout(len(regs), alt)
    data = [["", ""] * cols for _ in range(ROWS)]
    for i, (r, c) in enumerate(pos):
        data[r][2 * c:2 * c + 2] = [i + 1, regs[i]]
    return data, cols


def notice_rows(alloc, used_rooms):
    rows = []
    for room in used_rooms:
        sub = alloc[alloc.room_no == room]
        for (dept, code), g in sub.groupby(["department", "subcode"], sort=False):
            rows.append({
                "Dept": dept,
                "SubCode": code,
                "Regno": compress(g.reg_no),
                "Roomno": room,
                "Total": len(g),
            })
    return rows


# ---------------------------------------------------------------- excel
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def make_excel(alloc, used_rooms, title, date, session):
    rows = notice_rows(alloc, used_rooms)
    wb = Workbook()

    ws = wb.active
    ws.title = "Notice Board"
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", bold=True, size=14)
    ws["A2"] = f"DATE: {date}      SESSION: {session}"
    ws["A3"] = f"TOTAL NO OF ROOMS: {len(used_rooms)}      TOTAL NO OF STUDENTS: {len(alloc)}"
    for c in ("A2", "A3"):
        ws[c].font = Font(name="Arial", bold=True)
    head = ["Dept", "SubCode", "Regno", "Roomno", "Total"]
    for j, h in enumerate(head, 1):
        c = ws.cell(row=5, column=j, value=h)
        c.font = Font(name="Arial", bold=True)
        c.border = BOX
    for i, r in enumerate(rows, 6):
        for j, k in enumerate(head, 1):
            c = ws.cell(row=i, column=j, value=r[k])
            c.font = Font(name="Arial")
            c.border = BOX
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDE", (10, 10, 70, 10, 8)):
        ws.column_dimensions[col].width = w

    wa = wb.create_sheet("Allocation")
    cols = ["reg_no", "name", "department", "subcode", "room_no"]
    cols = [c for c in cols if c in alloc.columns]
    wa.append(cols)
    for _, r in alloc.iterrows():
        wa.append([r[c] for c in cols])
    for c in wa[1]:
        c.font = Font(name="Arial", bold=True)
    for col in "ABCDE":
        wa.column_dimensions[col].width = 18

    for room in used_rooms:
        w = wb.create_sheet(str(room)[:31])
        w["A1"] = f"ROOM: {room}"
        w["A1"].font = Font(name="Arial", bold=True, size=14)
        w["A2"] = f"DATE: {date}      SESSION: {session}"
        w.append([])
        sub_r = alloc[alloc.room_no == room]
        data, cols = room_grid(list(sub_r.reg_no), is_alt_room(sub_r))
        w.append(["S.No", "Reg No"] * cols)
        for c in w[4]:
            c.font = Font(name="Arial", bold=True)
            c.border = BOX
        for line in data:
            w.append(line)
            for c in w[w.max_row]:
                c.border = BOX
        for k in range(cols):
            w.column_dimensions[chr(65 + 2 * k)].width = 7
            w.column_dimensions[chr(66 + 2 * k)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- pdf
NOTICE_FS = 12  # notice board PDF font size (increase / decrease here)


def _table_style():
    return TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 0), (-1, -1), NOTICE_FS),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ])


def make_notice_pdf(alloc, used_rooms, title, date, session):
    ss = getSampleStyleSheet()
    small = ss["BodyText"]
    small.fontSize = NOTICE_FS
    small.leading = NOTICE_FS + 3
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=25, rightMargin=25, topMargin=25, bottomMargin=25)
    el = [
        Paragraph(f"<b>{title}</b>", ss["Title"]),
        Paragraph(f"<b>DATE:</b> {date} &nbsp;&nbsp; <b>SESSION:</b> {session}", ss["Normal"]),
        Paragraph(f"<b>TOTAL ROOMS:</b> {len(used_rooms)} &nbsp;&nbsp; <b>TOTAL STUDENTS:</b> {len(alloc)}", ss["Normal"]),
        Spacer(1, 10),
    ]
    data = [["Dept", "SubCode", "Regno", "Roomno", "Total"]]
    for r in notice_rows(alloc, used_rooms):
        data.append([r["Dept"], r["SubCode"], Paragraph(r["Regno"], small), r["Roomno"], r["Total"]])
    t = Table(data, colWidths=[58, 70, 290, 70, 55], repeatRows=1)
    t.setStyle(_table_style())
    el.append(t)
    doc.build(el)
    return buf.getvalue()


def make_room_pdf(alloc, used_rooms, title, date, session, center, caps, sign_path=None):
    """One A4 portrait page per room, in the university hall-chart format:
    header block, Subcode/Total strip, then Stno | Dept(+subcode) | Regno grid
    (5 rows, seats numbered column-wise), signature line at the bottom."""
    ss = getSampleStyleSheet()
    margin = 28
    width = A4[0] - 2 * margin

    title_st = ParagraphStyle("t", parent=ss["Normal"], fontName="Helvetica", fontSize=10, alignment=TA_CENTER)
    info_st = ParagraphStyle("i", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=8, leading=10)
    dept_st = ParagraphStyle("d", parent=ss["Normal"], fontName="Helvetica", alignment=TA_CENTER,
                             fontSize=7, leading=8)
    foot_st = ParagraphStyle("f", parent=ss["Normal"], fontName="Helvetica", fontSize=9)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=margin, rightMargin=margin,
                            topMargin=40, bottomMargin=30)
    el = []
    for k, room in enumerate(used_rooms):
        sub = alloc[alloc.room_no == room]
        regs = list(sub.reg_no)
        depts = list(sub.department)
        codes = list(sub.subcode)
        cap = caps.get(room, len(regs))

        el.append(Paragraph(escape(str(title)), title_st))
        el.append(Spacer(1, 22))

        # header block
        P = lambda s: Paragraph(escape(str(s)), info_st)
        info = [
            [P("EXAM CENTER CODE & NAME"), P(f":{center}"), P("DATE"), P(f":{date}")],
            [P("TOTAL SEATING CAPACITY"), P(f":{cap}"), P("SESSION"), P(f":{session}")],
            [P("HALL NO"), P(f":{room}"), P("HALL STRENGTH"), P(f":{len(regs)}")],
        ]
        t = Table(info, colWidths=[width * 0.30, width * 0.34, width * 0.18, width * 0.18])
        t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("TOPPADDING", (0, 0), (-1, -1), 2),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                               ("LEFTPADDING", (0, 0), (-1, -1), 2)]))
        el += [t, Spacer(1, 14)]

        # Subcode / Total strip (one pair per department in this room)
        strip_head, strip_val = [], []
        for (dept, code), g in sub.groupby(["department", "subcode"], sort=False):
            strip_head += ["Subcode", "Total"]
            strip_val += [code, len(g)]
        n = len(strip_head) // 2
        t = Table([strip_head, strip_val], colWidths=[62, 40] * n, hAlign="CENTER")
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.75, colors.black),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        el += [t, Spacer(1, 12)]

        # seat grid: Stno | Dept (+subcode) | Regno, 5 rows, numbered column-wise
        cols, pos = seat_layout(len(regs), is_alt_room(sub))
        at = {p: i for i, p in enumerate(pos)}
        group_w = width / cols
        w_no, w_dept = group_w * 0.15, group_w * 0.28
        w_reg = group_w - w_no - w_dept
        longest = max(len(str(r)) for r in regs)
        fs = min(9.5, (w_reg - 2) / (longest * 0.556))  # Helvetica-Bold digit = 0.556 em

        head = ["Stno", "Dept", "Regno"] * cols
        data = [head]
        for r in range(ROWS):
            line = []
            for c in range(cols):
                i = at.get((r, c))
                if i is not None:
                    dept_cell = Paragraph(
                        f'<font size="8"><b>{escape(str(depts[i]))}</b></font><br/>'
                        f'<font size="6">{escape(str(codes[i]))}</font>', dept_st)
                    line += [i + 1, dept_cell, regs[i]]
                else:
                    line += ["", "", ""]
            data.append(line)
        t = Table(data, colWidths=[w_no, w_dept, w_reg] * cols, rowHeights=[24] + [30] * ROWS)
        style = [
            ("GRID", (0, 0), (-1, -1), 0.75, colors.black),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 8),
            ("FONTSIZE", (0, 1), (-1, -1), fs),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 1),
            ("RIGHTPADDING", (0, 0), (-1, -1), 1),
        ]
        for c in range(cols):  # Stno column a bit smaller
            style.append(("FONTSIZE", (3 * c, 0), (3 * c, 0), 6))
            style.append(("FONTSIZE", (3 * c, 1), (3 * c, -1), 8))
        t.setStyle(TableStyle(style))
        el += [t, Spacer(1, 40)]

        left_txt = Paragraph("Name Sign of Hall Superintendent(Name in Block Letters)", foot_st)
        if sign_path and Path(sign_path).exists():
            w0, h0 = PILImage.open(sign_path).size
            iw = 165
            ih = iw * h0 / w0
            foot = Table([[left_txt, RLImage(str(sign_path), width=iw, height=ih, mask="auto")]],
                         colWidths=[width * 0.62, width * 0.38])
            foot.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                      ("LEFTPADDING", (0, 0), (0, 0), 40),
                                      ("TOPPADDING", (0, 0), (0, 0), 0.53 * ih - 8),
                                      ("ALIGN", (1, 0), (1, 0), "CENTER")]))
        else:
            foot = Table([[left_txt, Paragraph("CHIEF SUPERINTENDENT", foot_st)]],
                         colWidths=[width * 0.62, width * 0.38])
            foot.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 40)]))
        el.append(foot)

        if k < len(used_rooms) - 1:
            el.append(PageBreak())
    doc.build(el)
    return buf.getvalue()


def make_qp_split_pdf(alloc, used_rooms, title, date, session):
    """Question paper split-up: one row per subject code, one column per room. Landscape, no signature."""
    margin = 30
    width = landscape(A4)[0] - 2 * margin
    ss = getSampleStyleSheet()
    title_st = ParagraphStyle("qt", parent=ss["Normal"], fontName="Helvetica", fontSize=11, alignment=TA_CENTER)
    info_st = ParagraphStyle("qi", parent=ss["Normal"], fontName="Helvetica", fontSize=10, leading=13)

    counts = {k: int(v) for k, v in alloc.groupby(["subcode", "room_no"]).size().items()}
    data = [["SubCode"] + [str(r) for r in used_rooms] + ["Total"]]
    for sc in sorted(alloc.subcode.unique()):
        row = [counts.get((sc, r), 0) or "" for r in used_rooms]
        data.append([sc] + row + [sum(c for c in row if c != "")])
    data.append(["Total"] + [int((alloc.room_no == r).sum()) for r in used_rooms] + [len(alloc)])

    n = len(used_rooms)
    first_w, last_w = 64, 46
    col_w = (width - first_w - last_w) / n
    fs = 10 if col_w >= 34 else (8 if col_w >= 24 else 6.5)

    P = lambda s: Paragraph(escape(str(s)), info_st)
    info = Table([[P("TOTAL NO OF ROOMS"), P(f":{n}"), P("DATE"), P(date)],
                  [P("TOTAL NO OF STUDENTS"), P(f":{len(alloc)}"), P("SESSION"), P(session)]],
                 colWidths=[width * 0.28, width * 0.32, width * 0.14, width * 0.26])
    info.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                              ("LEFTPADDING", (0, 0), (-1, -1), 2),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))

    t = Table(data, colWidths=[first_w] + [col_w] * n + [last_w], rowHeights=30, repeatRows=1)
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.75, colors.black),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -2), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), fs),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
    ]))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=margin, rightMargin=margin,
                            topMargin=35, bottomMargin=30)
    doc.build([Paragraph(escape(str(title)), title_st), Spacer(1, 16), info, Spacer(1, 10), t])
    return buf.getvalue()


# ---------------------------------------------------------------- PDF input
DEFAULT_SHORT = {"103": "CIVIL", "104": "CSE", "105": "EEE", "106": "ECE",
                 "114": "MECH", "205": "IT", "243": "AI&DS"}


def auto_short(name):
    """'B.Tech. Information Technology' -> 'IT' (used only when branch code is not in DEFAULT_SHORT)"""
    name = re.sub(r"^\s*(B\.E\.|B\.Tech\.|M\.E\.|M\.Tech\.)\s*", "", name, flags=re.I)
    words = [w for w in re.split(r"[\s&]+", name) if w and w.lower() not in ("and", "of", "in")]
    return "".join(w[0].upper() for w in words) or name


def parse_attendance_pdf(file):
    """Reads the university attendance-sheet PDF.
    Returns (students DataFrame, meta dict). Pages without a header (continuation pages)
    use the branch/subject/date/session of the previous page."""
    reader = PdfReader(file)
    meta = {"exam": "", "college_code": "", "college_name": ""}
    cur, rows = {}, []
    for page in reader.pages:
        text = page.extract_text() or ""
        m = re.search(r"UNIVERSITY EXAMINATIONS\s*-\s*(.+?Examination,?\s*\d{4})", text)
        if m and not meta["exam"]:
            meta["exam"] = re.sub(r"\s+", " ", m.group(1))
        d = re.search(r"Date of Examination\s*:\s*(\d{1,2}-[A-Za-z]{3}-\d{4})", text)
        s_ = re.search(r"Session\s*:\s*(FN|AN)\b", text)
        b = re.search(r"Branch Code/Name\s*:\s*(\d+)\s*-\s*(.+?)\s+Session\s*:", text)
        sj = re.search(r"Subject Code/Name\s*:\s*([A-Za-z0-9]+)\s*-\s*(.+)", text)
        c = re.search(r"College Code/Name\s*:\s*(\d+)\s*-\s*(.+?)\s*Date of Examination", text, re.S)
        if c and not meta["college_code"]:
            meta["college_code"] = c.group(1)
            meta["college_name"] = re.sub(r"\s+", " ", c.group(2)).strip()
        if d and s_ and b and sj:
            cur = {"date": d.group(1), "session": s_.group(1),
                   "branch_code": b.group(1), "branch_name": b.group(2).strip(),
                   "subcode": sj.group(1), "subname": sj.group(2).strip()}
        for line in text.splitlines():
            m = re.match(r"^\s*(\d{1,3})\s+(\d{10,14})\s+(.+?)\s*$", line)
            if m:
                if not cur:
                    raise ValueError("PDF la branch/subject header kedaikala. Correct attendance sheet PDF ah?")
                rows.append({"reg_no": m.group(2), "name": m.group(3), **cur})
    if not rows:
        raise ValueError("PDF la students kedaikala. Correct attendance sheet PDF ah?")
    return pd.DataFrame(rows), meta


def exam_title(exam):
    """'April / May Examination,2026' -> 'ANNA UNIVERSITY THEORY EXAMINATIONS APR-MAY-2026'"""
    m = re.search(r"([A-Za-z]+)\s*/\s*([A-Za-z]+)\s+Examination,?\s*(\d{4})", exam)
    if m:
        tail = f"{m.group(1)[:3]}-{m.group(2)[:3]}-{m.group(3)}".upper()
    else:
        tail = exam.upper()
    return f"ANNA UNIVERSITY THEORY EXAMINATIONS {tail}".strip()


# ---------------------------------------------------------------- UI
def main():
    import streamlit as st

    st.set_page_config(page_title="Exam Seating", layout="wide")
    st.title("Exam Room Allotment")

    f1 = st.file_uploader("Attendance sheet PDF (college upload pannadhu)", type=["pdf"])
    if not f1:
        st.info("PDF upload pannunga.")
        return

    sig = f"{f1.name}_{f1.size}"
    if st.session_state.get("pdf_sig") != sig:
        try:
            st.session_state["pdf"] = parse_attendance_pdf(io.BytesIO(f1.getvalue()))
        except Exception as e:
            st.error(f"PDF padikka mudiyala: {e}")
            return
        st.session_state["pdf_sig"] = sig
        st.session_state.pop("res", None)
    students, meta = st.session_state["pdf"]

    # date + session come from the PDF
    combos = students[["date", "session"]].drop_duplicates().values.tolist()
    if len(combos) > 1:
        labels = [f"{d}  |  {s}" for d, s in combos]
        pick = st.selectbox("PDF la onnukku mela date/session irukku. Edhu venum?", labels)
        date, session = combos[labels.index(pick)]
        students = students[(students.date == date) & (students.session == session)]
    else:
        date, session = combos[0]

    dup = students[students.reg_no.duplicated(keep=False)]
    if len(dup):
        st.warning("Intha reg no onnukku mela thadavai irukku (first onnu mattum edukren): "
                   + ", ".join(sorted(dup.reg_no.unique())))
        students = students.drop_duplicates("reg_no", keep="first")
    n = len(students)

    st.success(f"Date: {date}   |   Session: {session}   |   Students: {n}")

    c1, c2 = st.columns(2)
    title = c1.text_input("Exam title", exam_title(meta["exam"]), key=f"title_{sig}")
    center = c2.text_input("Exam center code & name",
                           f"{meta['college_code']}-{meta['college_name']}", key=f"center_{sig}")

    # signature + stamp for the room wise PDF (saved as signature.png next to app.py)
    with st.expander("Room wise PDF signature + stamp"):
        up = st.file_uploader("Pudhu signature image maathanum na upload pannunga (png / jpg)",
                              type=["png", "jpg", "jpeg"], key="sign_up")
        if up:
            try:
                SIGN_FILE.write_bytes(clean_signature(up.getvalue()))
            except Exception as e:
                st.error(f"Signature image padikka mudiyala: {e}")
        if SIGN_FILE.exists():
            st.image(str(SIGN_FILE), width=220)
        else:
            st.caption("signature.png illa. Room wise PDF la 'CHIEF SUPERINTENDENT' text mattum varum.")

    # short department names (from branch code)
    with st.expander("Department short names (venumna maathunga)"):
        shorts = {}
        branches = students[["branch_code", "branch_name"]].drop_duplicates().values.tolist()
        cols = st.columns(4)
        for i, (code, name) in enumerate(branches):
            shorts[code] = cols[i % 4].text_input(
                f"{code} - {name}", DEFAULT_SHORT.get(code) or auto_short(name), key=f"short_{sig}_{code}")
    students = students.assign(department=students.branch_code.map(shorts))

    st.subheader("PDF la irundhu eduthadhu")
    summary = (students.groupby(["department", "subcode", "subname"]).size()
               .reset_index(name="Students").rename(columns={"department": "Dept", "subcode": "SubCode",
                                                              "subname": "Subject"}))
    st.dataframe(summary, use_container_width=True, hide_index=True)

    # rooms: count = students / default capacity (editable), room names are remembered
    # for the next exam (saved in rooms_config.json next to app.py)
    st.subheader("Rooms")
    if "room_cfg" not in st.session_state:
        st.session_state["room_cfg"] = load_room_cfg()
    cfg = st.session_state["room_cfg"]  # snapshot at session start, used only for defaults
    saved_names = cfg["room_names"]

    r1, r2 = st.columns(2)
    def_cap = int(r2.number_input("Room ku default capacity", min_value=1, max_value=35,
                                  value=min(max(cfg["default_capacity"], 1), 35), step=1,
                                  key="def_cap"))
    auto_rooms = max(1, -(-n // def_cap))
    n_rooms = int(r1.number_input(f"Rooms (auto: {n} / {def_cap} = {auto_rooms}. Venumna maathunga)",
                                  min_value=1, max_value=100, value=auto_rooms, step=1,
                                  key=f"n_rooms_{sig}_{n}_{def_cap}"))
    rcols = st.columns(5)
    entries = []
    left = n  # students not yet given a seat; the last needed room defaults to what is left (e.g. 7)
    for i in range(n_rooms):
        cap_default = min(def_cap, left) if left > 0 else def_cap
        left -= cap_default
        with rcols[i % 5]:
            room = st.text_input(f"Room {i + 1}", saved_names[i] if i < len(saved_names) else "",
                                 key=f"room_{i}", placeholder="Room number type pannunga").strip()
            cap_i = int(st.number_input("Capacity", min_value=1, max_value=35, value=cap_default,
                                        step=1, key=f"cap_{i}_{def_cap}_{n}_{cap_default}",
                                        label_visibility="collapsed"))
        entries.append({"room": room, "cap": cap_i})

    # remember room names + default capacity for the next exam
    names_now = [e["room"] for e in entries]
    if any(names_now):
        on_disk = load_room_cfg()
        merged = names_now + on_disk["room_names"][len(names_now):]  # keep names of unused rooms
        if merged != on_disk["room_names"] or def_cap != on_disk["default_capacity"]:
            save_room_cfg({"default_capacity": def_cap, "room_names": merged})

    rooms = [e["room"] for e in entries]
    if "" in rooms:
        st.warning("Ella room box um fill pannunga (room number type pannunga).")
        return
    if len(set(rooms)) != len(rooms):
        st.error("Oru room number rendu thadavai irukku. Check pannunga.")
        return

    caps = {e["room"]: e["cap"] for e in entries}
    total_cap = sum(caps.values())
    if total_cap < n:
        st.error(f"Total capacity {total_cap} dhaan. {n} students ku podhala. "
                 "Rooms ah / capacity ah ekkunga.")
        return

    # how many rooms will actually be used (rooms are filled in the order given)
    run, need = 0, 0
    for r in rooms:
        run += caps[r]
        need += 1
        if run >= n:
            break
    last_room = rooms[need - 1]
    last = n - (run - caps[last_room])
    st.write(f"{n} students ku **{need} rooms** venum  ({n_rooms} rooms set, total capacity {total_cap}). "
             f"Last room ({last_room}) la {last} students.")

    if st.button("Allocate", type="primary"):
        try:
            alloc, used = allocate(students[["reg_no", "name", "department", "subcode"]],
                                   [(r, caps[r]) for r in rooms])
        except ValueError as e:
            st.error(str(e))
            return
        st.session_state["res"] = (alloc, used, caps)

    if "res" in st.session_state:
        alloc, used, caps = st.session_state["res"]
        st.success(f"{len(alloc)} students, {len(used)} rooms")
        st.dataframe(pd.DataFrame(notice_rows(alloc, used)), use_container_width=True)
        d1, d2, d3, d4 = st.columns(4)
        d1.download_button("Excel (all sheets)", make_excel(alloc, used, title, date, session), "allotment.xlsx")
        d2.download_button("Notice board PDF", make_notice_pdf(alloc, used, title, date, session), "notice_board.pdf")
        sign = SIGN_FILE if SIGN_FILE.exists() else None
        d3.download_button("Room wise PDF", make_room_pdf(alloc, used, title, date, session, center, caps, sign),
                           "room_wise.pdf")
        d4.download_button("QP split-up PDF (landscape)", make_qp_split_pdf(alloc, used, title, date, session),
                           f"QSplitUp_{date}_{session}.pdf")


if __name__ == "__main__":
    main()
